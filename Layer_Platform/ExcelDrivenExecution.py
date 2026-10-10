from openpyxl import Workbook
from openpyxl.reader.excel import load_workbook
import pywintypes
import win32com
import win32com.client as w32c
from win32com.client import gencache, Dispatch, DispatchWithEvents, constants
import xlsxwriter
import os
import openpyxl
import datetime
from openpyxl.worksheet.copier import WorksheetCopy
import pyodbc
import numpy as np


##  Find active modules for execution
def FindActiveModules(sModuleFilePath):
	try:
		wb=load_workbook(sModuleFilePath)
		ws=wb.get_sheet_by_name("tbl_modules")
		SearchString ="Yes"
		Modules = ""
		iRow = 1
		for row in ws.iter_rows():
			for cell in row:			
				if(str(SearchString)==str(cell.value)):
					Modules = Modules + str(ws.cell(column=1, row=iRow).value) + ".xlsx" + ":"
				
			iRow = iRow +1
				
		if Modules== "":
			Modules = "No modules selected"
		return Modules
	except Exception,e:
		err = str(e)
		print err
		return False
 

 ##  Find active testcases of the module for execution
def FindActiveTestCases(sTestCasesFilePath):
	print sTestCasesFilePath
	try:
		wb=load_workbook(sTestCasesFilePath)
		ws=wb.get_sheet_by_name("tbl_testcase")
		SearchString ="Yes"
		Testcases = ""
		TCIDs = ""
		iRow = 1
		TestCaseData = {}
		for row in ws.iter_rows():
			for cell in row:			
				if(str(SearchString)==str(cell.value)):
					Testcases = Testcases + str(ws.cell(column=2, row=iRow).value) + " "
					# TCIDs = TCIDs + str(ws.cell(column=3, row=iRow).value) + ".xlsx" + ":"
					# TestCaseData[iRow-2,0,0] = str(ws.cell(column=2, row=iRow).value)
					# TestCaseData[iRow-2,0,1] = str(ws.cell(column=3, row=iRow).value)
					# TestCaseData[iRow-2,1,1] = str(ws.cell(column=7, row=iRow).value)
				
			iRow = iRow +1
				
		if Testcases== "":
			Testcases = "No Testcases selected"
		return Testcases
	except Exception,e:
		err = str(e)
		print err
		return False
